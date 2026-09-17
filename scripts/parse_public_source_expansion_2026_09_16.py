#!/usr/bin/env python3
"""Normalize acquired state WARN notices with record-level provenance."""

from __future__ import annotations

import argparse
import calendar
import csv
import hashlib
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import load_workbook

AS_OF = "2026-09-16"
SOURCES = {
    "CA": {
        "filename": f"ca_warn_latest_{AS_OF}.xlsx",
        "url": "https://edd.ca.gov/siteassets/files/jobs_and_training/warn/warn_report1.xlsx",
        "scope": "California WARN reports processed July 1-September 14, 2026",
    },
    "TX": {
        "filename": f"tx_warn_2026_{AS_OF}.xlsx",
        "url": "https://www.twc.texas.gov/sites/default/files/oei/docs/warn-act-listings-2026-twc.xlsx",
        "scope": "Texas notices listed for calendar 2026 through September 15",
    },
    "PA": {
        "filename": f"pa_warn_2026_{AS_OF}.html",
        "url": "https://www.pa.gov/agencies/dli/programs-services/workforce-development-home/warn-requirements/warn-notices",
        "scope": "Pennsylvania 2026 public notices; month is the only notice timing field",
    },
}
FIELDS = [
    "state",
    "notice_date",
    "received_date",
    "notice_month",
    "effective_date_text",
    "company",
    "county",
    "affected_workers_reported",
    "action_text",
    "address",
    "industry",
    "source_uri",
    "source_record_ref",
    "content_hash",
    "retrieved_pacific_date",
]


def iso(value: object) -> str:
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return "" if value is None else str(value).strip()


def integer(value: object) -> int | None:
    if value is None:
        return None
    match = re.search(r"\d[\d,]*", str(value))
    return int(match.group().replace(",", "")) if match else None


def base_row(state: str, record_ref: str) -> dict[str, str | int | None]:
    return {
        **dict.fromkeys(FIELDS, ""),
        "state": state,
        "source_uri": SOURCES[state]["url"],
        "source_record_ref": record_ref,
        "retrieved_pacific_date": AS_OF,
    }


def parse_ca(path: Path) -> list[dict[str, str | int | None]]:
    sheet = load_workbook(path, read_only=True, data_only=True)["Detailed WARN Report "]
    rows = []
    for row_index, values in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
        if not values[4] or not str(values[4]).strip():
            continue
        row = base_row("CA", f"Detailed WARN Report !{row_index}")
        row.update(
            notice_date=iso(values[1]),
            received_date=iso(values[2]),
            effective_date_text=iso(values[3]),
            company=str(values[4]).strip(),
            county=str(values[0] or "").strip(),
            action_text=str(values[5] or "").strip(),
            affected_workers_reported=integer(values[6]),
            address=str(values[7] or "").strip(),
            industry=str(values[8] or "").strip(),
        )
        rows.append(row)
    return rows


def parse_tx(path: Path) -> list[dict[str, str | int | None]]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = []
    for row_index, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if not values[1] or not str(values[1]).strip():
            continue
        row = base_row("TX", f"Sheet1!{row_index}")
        row.update(
            notice_date=iso(values[0]),
            received_date=iso(values[6]),
            effective_date_text=iso(values[5]),
            company=str(values[1]).strip(),
            county=str(values[2] or "").strip(),
            affected_workers_reported=integer(values[4]),
            address=str(values[7] or "").strip(),
        )
        rows.append(row)
    return rows


def _label(text: str, label: str, next_label: str | None = None) -> str:
    end = rf"(?=\s*{next_label})" if next_label else "$"
    match = re.search(rf"{label}\s*:\s*(.*?){end}", text, re.I)
    return match.group(1).strip() if match else ""


def parse_pa(path: Path) -> list[dict[str, str | int | None]]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    heading = soup.find("h2", string="2026")
    if heading is None:
        raise ValueError("Pennsylvania 2026 accordion not found")
    rows = []
    for month_item in heading.parent.find_all("div", class_="cmp-accordion__item", recursive=False):
        month = month_item.find("span", class_="cmp-accordion__title").get_text(" ", strip=True)
        month_panel = month_item.find("div", class_="cmp-accordion__panel", recursive=False)
        for notice in month_panel.select("div.cmp-accordion__item"):
            title = notice.find("h3", class_="cmp-accordion__header")
            if title is None:
                continue
            name = title.find("span", class_="cmp-accordion__title")
            if name is None:
                continue
            panel = notice.find("div", class_="cmp-accordion__panel", recursive=False)
            if panel is None:
                continue
            text = panel.get_text(" ", strip=True).replace("\u200b", "")
            affected = integer(_label(text, r"#\s*AFFECTED", "EFFECTIVE DATE"))
            row = base_row("PA", notice.get("id") or "")
            row.update(
                notice_month=f"2026-{list(calendar.month_name).index(month):02d}",
                effective_date_text=_label(text, "EFFECTIVE DATE", "CLOSURE OR LAYOFF"),
                company=name.get_text(" ", strip=True),
                county=_label(text, "COUNTY", r"#\s*AFFECTED"),
                affected_workers_reported=affected,
                action_text=_label(text, "CLOSURE OR LAYOFF"),
                address=text.split("COUNTY:", 1)[0].strip(),
            )
            rows.append(row)
    return rows


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    parse = {"CA": parse_ca, "TX": parse_tx, "PA": parse_pa}
    rows = []
    sources = {}
    for state, source in SOURCES.items():
        path = args.raw_dir / source["filename"]
        if not path.is_file():
            raise FileNotFoundError(path)
        current = parse[state](path)
        raw_hash = sha256(path)
        for row in current:
            row["content_hash"] = raw_hash
        rows.extend(current)
        sources[state] = {
            "url": source["url"],
            "scope": source["scope"],
            "raw_path": str(path),
            "raw_sha256": raw_hash,
            "record_count": len(current),
            "reported_affected_sum": sum(r["affected_workers_reported"] or 0 for r in current),
        }
    output = args.output_dir / "warn_notices.csv"
    with output.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    summary = {
        "as_of_pacific_date": AS_OF,
        "normalized_path": str(output),
        "normalized_sha256": sha256(output),
        "total_notice_rows": len(rows),
        "states": sources,
        "notice_rows_by_state": dict(Counter(r["state"] for r in rows)),
        "caveats": [
            "WARN notices are announced planned layoffs/closures, not confirmed separations.",
            "State WARN legal thresholds and publication practices differ; totals are not comparable rates.",
            "Pennsylvania gives month-level notice timing; its revised notices can overlap earlier entries.",
            "The three-state sample cannot establish national layoffs or AI-caused layoffs.",
        ],
    }
    summary_path = args.output_dir / "warn_summary.json"
    summary_path.write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n")
    print(json.dumps(summary, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
