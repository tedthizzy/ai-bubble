from io import BytesIO

import pytest
from openpyxl import Workbook

from bubble.ingestion.sources.catalog import _parse_xlsx


def test_required_table_header_after_thirty_rows_keeps_all_projects():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Project Details - Large Gen"
    for _ in range(30):
        sheet.append(["Published notes", None, None])
    sheet.append(["INR", "Project Name", "Capacity (MW)"])
    sheet.append(["A1", "Project one", 100])
    sheet.append(["A2", "Project two", 200])
    small = workbook.create_sheet("Project Details - Small Gen")
    small.append(["INR", "Project Name", "Capacity (MW)"])
    small.append(["B1", "Project three", 3])
    stream = BytesIO()
    workbook.save(stream)
    rows = _parse_xlsx(
        stream.getvalue(),
        metadata={
            "xlsx_required_value_columns": "INR|Project Name",
            "xlsx_sheet_name_prefix": "Project Details -",
        },
    )
    assert [row["INR"] for row in rows] == ["A1", "A2", "B1"]
    assert [row["source_row_number"] for row in rows] == [32, 33, 2]


def test_missing_declared_header_is_an_error():
    workbook = Workbook()
    workbook.active.append(["Unrelated", "Table"])
    stream = BytesIO()
    workbook.save(stream)
    with pytest.raises(ValueError, match="required columns"):
        _parse_xlsx(stream.getvalue(), metadata={"xlsx_required_value_columns": "INR"})
